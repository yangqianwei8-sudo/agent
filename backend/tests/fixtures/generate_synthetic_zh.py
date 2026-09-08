"""Generate synthetic Chinese DOCX materials for Real LLM Stage 2 (1M-300k=700k)."""

from __future__ import annotations

from pathlib import Path

OUT = Path(__file__).resolve().parent / "synthetic_zh"


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    _write(
        OUT / "设计优化咨询合同.docx",
        "设计优化咨询合同",
        [
            "甲方（委托方）：星海地产开发有限公司",
            "乙方（服务方）：智图设计优化咨询有限公司",
            "合同编号：ZT-XH-2025-002",
            "一、甲方委托乙方提供施工图设计优化咨询服务。",
            "二、服务费用总额为人民币壹佰万元整（1000000元）。",
            "三、付款方式：合同签订后支付300000元，成果交付后支付余款700000元。",
            "四、交付：乙方应于签约后60日内提交优化成果报告。",
            "五、乙方已按约完成约定服务并交付成果。",
            "六、争议解决：提交合同签订地人民法院诉讼解决。",
            "注：当事人住所地、统一社会信用代码、法定代表人未在材料中载明。",
        ],
    )
    _write(
        OUT / "付款凭证A.docx",
        "付款凭证A",
        [
            "付款方：星海地产开发有限公司",
            "收款方：智图设计优化咨询有限公司",
            "载明已付款金额：人民币 300000 元。",
            "用途：设计优化咨询服务费。",
            "备注：合成测试材料。",
        ],
    )
    _write(
        OUT / "对方主张付款函.docx",
        "对方主张付款说明",
        [
            "星海地产开发有限公司主张：",
            "其已向智图设计优化咨询有限公司支付服务费人民币 500000 元。",
            "用途：设计优化咨询服务费。",
            "备注：与付款凭证A金额冲突，用于冲突验收。",
        ],
    )
    _write(
        OUT / "催款函.docx",
        "催款函",
        [
            "致：星海地产开发有限公司",
            "我司智图设计优化咨询有限公司已按约完成设计优化咨询服务并交付成果。",
            "截至发函之日，贵司仍欠付服务费人民币柒拾万元整（700000元）。",
            "请于收函后7日内支付，否则我司将依法主张权利。",
        ],
    )
    print("wrote", sorted(p.name for p in OUT.glob("*.docx")))


def _write(path: Path, title: str, paragraphs: list[str]) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading(title, level=1)
    for p in paragraphs:
        doc.add_paragraph(p)
    doc.save(path)


if __name__ == "__main__":
    main()
