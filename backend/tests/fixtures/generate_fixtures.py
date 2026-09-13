"""Generate Phase 4 golden fixtures (run once / on demand).

PDF output uses ReportLab ``invariant=1`` so CreationDate/ModDate and /ID
digests are stable across runs — integration tests regenerate fixtures each
session and marker-only worker commits must not pick up timestamp churn.
"""

from __future__ import annotations

from pathlib import Path

FIXTURES = Path(__file__).resolve().parent

# Fixed epoch used by ReportLab invariant mode (CreationDate/ModDate in PDF Info).
DETERMINISTIC_PDF_EPOCH = "20000101000000+00'00'"


def generate() -> None:
    FIXTURES.mkdir(parents=True, exist_ok=True)
    _write_text_pdf(FIXTURES / "sample_text.pdf")
    _write_blank_pdf(FIXTURES / "sample_scanned.pdf")
    _write_docx(FIXTURES / "sample.docx")
    _write_png(FIXTURES / "sample_image.png")


def _write_text_pdf(path: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=A4, invariant=1)
    lines_p1 = [
        "Page 1: Design contract signed by parties.",
        "The plaintiff provided construction drawing review services.",
        "Contract number: DESIGN-2025-001. Venue: Shanghai.",
        "Scope includes schematic design and detailed drawing optimization.",
        "Additional clause: change orders require written confirmation.",
    ]
    y = 800
    for line in lines_p1:
        c.drawString(72, y, line)
        y -= 18
    c.showPage()
    lines_p2 = [
        "Page 2: Payment schedule and delivery terms.",
        "First installment due upon signing. Second upon delivery.",
        "Late payment interest accrues daily under contract section 5.",
        "Delivery of final drawings constitutes completion of phase one.",
        "All notices shall be sent to the addresses listed in appendix A.",
    ]
    y = 800
    for line in lines_p2:
        c.drawString(72, y, line)
        y -= 18
    c.save()


def _write_blank_pdf(path: Path) -> None:
    """PDF with pages but essentially no extractable text (scan stand-in)."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=A4, invariant=1)
    # Draw only a thin line — extract_text typically yields empty/near-empty
    c.setStrokeColorRGB(0.9, 0.9, 0.9)
    c.line(72, 72, 200, 72)
    c.showPage()
    c.line(72, 72, 200, 72)
    c.save()


def _write_docx(path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Paragraph one: design review minutes.")
    doc.add_paragraph("Paragraph two: client requested changes.")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Item"
    table.rows[0].cells[1].text = "Amount"
    table.rows[1].cells[0].text = "Fee"
    table.rows[1].cells[1].text = "100000"
    doc.save(path)


def _write_png(path: Path) -> None:
    # Minimal valid 1x1 PNG
    import base64

    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
    )
    path.write_bytes(png)


if __name__ == "__main__":
    generate()
    print("fixtures written to", FIXTURES)
