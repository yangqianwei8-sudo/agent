"""DOCX text extraction tool — no Repository / Domain writes."""

from __future__ import annotations

from io import BytesIO

from docx import Document

from backend.tools.dto import ParseFailureDTO, ParseResultDTO, ParseSuccessDTO, SpanProposal
from backend.tools.errors import UnsupportedFormatError


class DocxParser:
    EXTRACTION_METHOD = "python_docx"
    EXTRACTION_VERSION = "v1"

    def parse(self, data: bytes, *, filename: str | None = None) -> ParseResultDTO:
        name = (filename or "").lower()
        if name.endswith(".doc") and not name.endswith(".docx"):
            raise UnsupportedFormatError(".doc is not supported in V1; convert to .docx")

        try:
            document = Document(BytesIO(data))
        except Exception as exc:  # noqa: BLE001
            return ParseFailureDTO(
                error_code="DOCX_OPEN_FAILED",
                error_detail=str(exc)[:500],
                extraction_method=self.EXTRACTION_METHOD,
                extraction_version=self.EXTRACTION_VERSION,
            )

        blocks: list[tuple[int | None, str]] = []
        para_index = 0

        for paragraph in document.paragraphs:
            text = paragraph.text or ""
            para_index += 1
            blocks.append((para_index, text))

        for table in document.tables:
            for row in table.rows:
                cells = [(cell.text or "").strip() for cell in row.cells]
                line = " | ".join(cells)
                para_index += 1
                blocks.append((para_index, line))

        if not blocks:
            return ParseFailureDTO(
                error_code="DOCX_EMPTY",
                error_detail="document has no paragraphs or tables",
                extraction_method=self.EXTRACTION_METHOD,
                extraction_version=self.EXTRACTION_VERSION,
            )

        separator = "\n"
        spans: list[SpanProposal] = []
        full_parts: list[str] = []
        cursor = 0
        paragraph_map: list[dict[str, int]] = []
        for i, (para_no, text) in enumerate(blocks):
            start = cursor
            end = start + len(text)
            if para_no is not None:
                paragraph_map.append({"paragraph": para_no, "start": start, "end": end})
            if end > start:
                spans.append(
                    SpanProposal(
                        character_start=start,
                        character_end=end,
                        quote=text,
                        page=None,
                        paragraph=para_no,
                        bbox_json=None,
                        confidence=1.0,
                        weak_localization=False,
                    )
                )
            full_parts.append(text)
            cursor = end
            if i < len(blocks) - 1:
                cursor += len(separator)

        full_text = separator.join(full_parts)
        layout = {
            "page_map": [],
            "paragraph_map": paragraph_map,
            "page_break_policy": "none",
            "separator": separator,
        }

        if not spans:
            return ParseFailureDTO(
                error_code="DOCX_NO_TEXT_SPANS",
                error_detail="no extractable text spans",
                extraction_method=self.EXTRACTION_METHOD,
                extraction_version=self.EXTRACTION_VERSION,
            )

        return ParseSuccessDTO(
            full_text=full_text,
            page_count=None,
            layout_json=layout,
            spans=spans,
            extraction_method=self.EXTRACTION_METHOD,
            extraction_version=self.EXTRACTION_VERSION,
        )
