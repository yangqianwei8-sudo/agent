"""Production deterministic fixture generator for Phase 4 golden files.

PDF output uses ReportLab ``invariant=1`` so CreationDate/ModDate and /ID
digests are stable across runs.  Integration tests call ``ensure_fixtures()``
so committed golden files are not rewritten each session and marker-only
worker commits do not pick up timestamp churn.
"""

from __future__ import annotations

import re
import tempfile
from pathlib import Path

DEFAULT_FIXTURES_DIR = Path(__file__).resolve().parents[1] / "tests" / "fixtures"

# Fixed epoch used by ReportLab invariant mode (CreationDate/ModDate in PDF Info).
DETERMINISTIC_PDF_EPOCH = "20000101000000+00'00'"

FIXTURE_NAMES = ("sample_text.pdf", "sample_scanned.pdf", "sample.docx", "sample_image.png")

_PDF_DATE_RE = re.compile(r"(CreationDate|ModDate) \(D:([^)]+)\)")
_PDF_ID_RE = re.compile(r"/ID\s*\[<([0-9a-f]+)><\1>\]")


def pdf_has_deterministic_metadata(pdf_bytes: bytes) -> bool:
    """Return True when PDF bytes carry pinned CreationDate/ModDate and stable /ID."""
    text = pdf_bytes.decode("latin-1")
    dates = _PDF_DATE_RE.findall(text)
    if not dates:
        return False
    if any(value != DETERMINISTIC_PDF_EPOCH for _, value in dates):
        return False
    return _PDF_ID_RE.search(text) is not None


def validate_fixtures(root: Path) -> None:
    """Raise ValueError when committed fixtures drift from deterministic generation."""
    missing = [name for name in FIXTURE_NAMES if not (root / name).is_file()]
    if missing:
        raise ValueError(f"missing fixtures: {', '.join(missing)}")
    with tempfile.TemporaryDirectory() as tmp:
        expected = generate(output_dir=Path(tmp))
        drifted = [
            name
            for name in FIXTURE_NAMES
            if (root / name).read_bytes() != (expected / name).read_bytes()
        ]
    if drifted:
        raise ValueError(f"fixture drift detected: {', '.join(drifted)}")


def ensure_fixtures(*, output_dir: Path | None = None) -> Path:
    """Return fixture directory, generating when missing and validating when present."""
    root = output_dir if output_dir is not None else DEFAULT_FIXTURES_DIR
    if all((root / name).is_file() for name in FIXTURE_NAMES):
        validate_fixtures(root)
        return root
    return generate(output_dir=output_dir)


def generate(*, output_dir: Path | None = None) -> Path:
    """Write golden fixtures. Returns the output directory (committed or temp)."""
    root = output_dir if output_dir is not None else DEFAULT_FIXTURES_DIR
    root.mkdir(parents=True, exist_ok=True)
    _write_text_pdf(root / "sample_text.pdf")
    _write_blank_pdf(root / "sample_scanned.pdf")
    _write_docx(root / "sample.docx")
    _write_png(root / "sample_image.png")
    return root


def _write_text_pdf(path: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=A4, invariant=1)
    lines_p1 = [
        "Page 1: Design contract signed by parties.",
        "The plaintiff provided construction drawing review services.",
        "Contract number: DESIGN-2025-001. Venue: Shanghai.",
        "Scope includes schematic design and detailed drawing optimization.",
        "Additional clause: change orders require written confirmation.",
    ]
    y = 800
    for line in lines_p1:
        c.drawString(72, y, line)
        y -= 18
    c.showPage()
    lines_p2 = [
        "Page 2: Payment schedule and delivery terms.",
        "First installment due upon signing. Second upon delivery.",
        "Late payment interest accrues daily under contract section 5.",
        "Delivery of final drawings constitutes completion of phase one.",
        "All notices shall be sent to the addresses listed in appendix A.",
    ]
    y = 800
    for line in lines_p2:
        c.drawString(72, y, line)
        y -= 18
    c.save()


def _write_blank_pdf(path: Path) -> None:
    """PDF with pages but essentially no extractable text (scan stand-in)."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=A4, invariant=1)
    # Draw only a thin line — extract_text typically yields empty/near-empty
    c.setStrokeColorRGB(0.9, 0.9, 0.9)
    c.line(72, 72, 200, 72)
    c.showPage()
    c.line(72, 72, 200, 72)
    c.save()


def _write_docx(path: Path) -> None:
    import io
    import zipfile

    from docx import Document

    doc = Document()
    doc.add_paragraph("Paragraph one: design review minutes.")
    doc.add_paragraph("Paragraph two: client requested changes.")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Item"
    table.rows[0].cells[1].text = "Amount"
    table.rows[1].cells[0].text = "Fee"
    table.rows[1].cells[1].text = "100000"
    raw = io.BytesIO()
    doc.save(raw)
    raw.seek(0)
    fixed_zip_dt = (2000, 1, 1, 0, 0, 0)
    with zipfile.ZipFile(raw, "r") as src, zipfile.ZipFile(path, "w") as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            normalized = zipfile.ZipInfo(filename=info.filename)
            normalized.compress_type = info.compress_type
            normalized.external_attr = info.external_attr
            normalized.date_time = fixed_zip_dt
            dst.writestr(normalized, data)


def _write_png(path: Path) -> None:
    # Minimal valid 1x1 PNG
    import base64

    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
    )
    path.write_bytes(png)
